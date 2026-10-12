#!/usr/bin/env python3
"""
SmartTrack — Product / Development Requirements Deck
PM-style documentation slides with design placeholders.
"""

from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).parent
REF = ROOT / "ppt_ref.pptx"
OUT = ROOT / "DOKUMENTASI-SMARTTRACK-v0.3.pptx"

# ── Brand palette ───────────────────────────────────────────────────
NAVY = RGBColor(0x00, 0x20, 0x60)
BLUE = RGBColor(0x15, 0x65, 0xD8)
TEAL = RGBColor(0x00, 0x97, 0xA7)
DARK = RGBColor(0x1A, 0x27, 0x44)
GRAY = RGBColor(0x5C, 0x6B, 0x88)
LIGHT_BG = RGBColor(0xF4, 0xF6, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PLACEHOLDER_BG = RGBColor(0xE8, 0xEC, 0xF4)
PLACEHOLDER_BORDER = RGBColor(0xBD, 0xC5, 0xD8)
ORANGE_LABEL = RGBColor(0xEF, 0x6C, 0x00)
GREEN_LABEL = RGBColor(0x2E, 0x7D, 0x32)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN_L = Inches(0.55)
MARGIN_R = Inches(0.55)
CONTENT_W = SLIDE_W - MARGIN_L - MARGIN_R
FOOTER_Y = Inches(7.05)


class Deck:
    def __init__(self):
        shutil.copy(REF, OUT)
        self.prs = Presentation(str(OUT))
        # Remove default slides — rebuild from scratch keeping theme
        while len(self.prs.slides) > 0:
            self._delete_slide(0)
        self._slide_no = 0

    def _delete_slide(self, index: int) -> None:
        sld_id = self.prs.slides._sldIdLst[index]
        self.prs.part.drop_rel(sld_id.rId)
        del self.prs.slides._sldIdLst[index]

    def _blank(self):
        self._slide_no += 1
        return self.prs.slides.add_slide(self.prs.slide_layouts[14])  # Blank, no page

    def _footer(self, slide, section: str = "") -> None:
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, FOOTER_Y, SLIDE_W, Inches(0.45)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = LIGHT_BG
        bar.line.fill.background()
        tx = slide.shapes.add_textbox(MARGIN_L, FOOTER_Y + Inches(0.08), Inches(8), Inches(0.3))
        tf = tx.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = f"SmartTrack · Smart Security Dashboard  |  Draft v0.3  |  {section}"
        r.font.size = Pt(8)
        r.font.color.rgb = GRAY
        pg = slide.shapes.add_textbox(SLIDE_W - Inches(1.1), FOOTER_Y + Inches(0.08), Inches(0.5), Inches(0.3))
        p2 = pg.text_frame.paragraphs[0]
        r2 = p2.add_run()
        r2.text = str(self._slide_no)
        r2.font.size = Pt(8)
        r2.font.color.rgb = GRAY

    def _title_bar(self, slide, title: str, subtitle: str = "") -> float:
        """Navy header band. Returns Y below header."""
        h = Inches(1.05) if subtitle else Inches(0.85)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, h)
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        tx = slide.shapes.add_textbox(MARGIN_L, Inches(0.18), CONTENT_W, Inches(0.5))
        tf = tx.text_frame
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = title
        r.font.size = Pt(22)
        r.font.bold = True
        r.font.color.rgb = WHITE
        if subtitle:
            tx2 = slide.shapes.add_textbox(MARGIN_L, Inches(0.62), CONTENT_W, Inches(0.35))
            p2 = tx2.text_frame.paragraphs[0]
            r2 = p2.add_run()
            r2.text = subtitle
            r2.font.size = Pt(11)
            r2.font.color.rgb = RGBColor(0xBB, 0xCC, 0xEE)
        return h + Inches(0.2)

    def _section_divider(self, number: str, title: str, subtitle: str) -> None:
        slide = self._blank()
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        num = slide.shapes.add_textbox(MARGIN_L, Inches(2.2), Inches(2), Inches(1))
        p = num.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = number
        r.font.size = Pt(48)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x40, 0x60, 0xA0)
        tx = slide.shapes.add_textbox(MARGIN_L, Inches(3.1), Inches(10), Inches(1.2))
        p2 = tx.text_frame.paragraphs[0]
        r2 = p2.add_run()
        r2.text = title
        r2.font.size = Pt(32)
        r2.font.bold = True
        r2.font.color.rgb = WHITE
        tx3 = slide.shapes.add_textbox(MARGIN_L, Inches(4.2), Inches(10), Inches(0.6))
        p3 = tx3.text_frame.paragraphs[0]
        r3 = p3.add_run()
        r3.text = subtitle
        r3.font.size = Pt(14)
        r3.font.color.rgb = RGBColor(0xBB, 0xCC, 0xEE)
        self._footer(slide, title)

    def _label_chip(self, slide, text: str, left, top, color: RGBColor) -> None:
        w = Inches(len(text) * 0.09 + 0.35)
        chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, Inches(0.28))
        chip.fill.solid()
        chip.fill.fore_color.rgb = color
        chip.line.fill.background()
        tf = chip.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = WHITE

    def _text_block(self, slide, left, top, width, height, lines: list[str], *, size=Pt(11), bold_first=False):
        tx = slide.shapes.add_textbox(left, top, width, height)
        tf = tx.text_frame
        tf.word_wrap = True
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            if line.startswith("## "):
                p.text = line[3:]
                for run in p.runs:
                    run.font.bold = True
                    run.font.size = Pt(12)
                    run.font.color.rgb = NAVY
                p.space_before = Pt(8)
            elif line.startswith("—"):
                p.text = line
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = GRAY
                    run.font.italic = True
            elif line == "":
                p.text = ""
            else:
                p.text = line
                for run in p.runs:
                    run.font.size = size
                    run.font.color.rgb = DARK
                    if bold_first and i == 0:
                        run.font.bold = True
            p.space_after = Pt(3)
        return tx

    def _screenshot_placeholder(
        self, slide, left, top, width, height, label: str, caption: str = ""
    ) -> None:
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        box.fill.solid()
        box.fill.fore_color.rgb = PLACEHOLDER_BG
        box.line.color.rgb = PLACEHOLDER_BORDER
        box.line.width = Pt(1.5)
        # dashed effect via inner text
        cx = left + width / 2 - Inches(2.5)
        cy = top + height / 2 - Inches(0.55)
        tx = slide.shapes.add_textbox(cx, cy, Inches(5), Inches(1.1))
        tf = tx.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = "[ INSERT SCREENSHOT ]"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = GRAY
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = label
        r2.font.size = Pt(10)
        r2.font.color.rgb = BLUE
        if caption:
            p3 = tf.add_paragraph()
            p3.alignment = PP_ALIGN.CENTER
            r3 = p3.add_run()
            r3.text = caption
            r3.font.size = Pt(9)
            r3.font.color.rgb = GRAY

    def _diagram_placeholder(self, slide, left, top, width, height, label: str) -> None:
        self._screenshot_placeholder(slide, left, top, width, height, label, "Diagram / flowchart — dapat digambar di Visio")

    # ── Slide types ───────────────────────────────────────────────────

    def cover(self) -> None:
        slide = self._blank()
        # left panel
        panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(5.2), SLIDE_H)
        panel.fill.solid()
        panel.fill.fore_color.rgb = NAVY
        panel.line.fill.background()

        lines = [
            ("SmartTrack", Pt(34), True, WHITE),
            ("Smart Security Dashboard", Pt(18), True, RGBColor(0xBB, 0xCC, 0xEE)),
            ("", Pt(11), False, WHITE),
            ("Product & Development", Pt(13), False, RGBColor(0xBB, 0xCC, 0xEE)),
            ("Requirements Document", Pt(13), False, RGBColor(0xBB, 0xCC, 0xEE)),
        ]
        y = Inches(1.8)
        for text, size, bold, color in lines:
            if text:
                tx = slide.shapes.add_textbox(Inches(0.55), y, Inches(4.5), Inches(0.6))
                p = tx.text_frame.paragraphs[0]
                r = p.add_run()
                r.text = text
                r.font.size = size
                r.font.bold = bold
                r.font.color.rgb = color
            y += Inches(0.55 if text else 0.2)

        meta = slide.shapes.add_textbox(Inches(0.55), Inches(5.6), Inches(4.5), Inches(1.2))
        tf = meta.text_frame
        for i, line in enumerate([
            "Versi 0.3  ·  Juni 2026",
            "OPTR · NTD Development",
            "Sinarmas Agribusiness and Food",
            "Status: Draft — Review Bisnis & IT",
        ]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = line
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0x99, 0xAA, 0xCC)

        # right — doc purpose
        self._label_chip(slide, "PRD / FSD", Inches(5.6), Inches(0.7), BLUE)
        self._text_block(slide, Inches(5.6), Inches(1.15), Inches(7.2), Inches(5.5), [
            "## Tujuan Dokumen",
            "Dokumen ini menjadi acuan awal bagi tim IT dalam mengembangkan modul patroli berbasis GPS di Smart Security Dashboard.",
            "",
            "## Audiens",
            "• Tim IT Development & Arsitektur",
            "• Operasional Keamanan Estate",
            "• PIC & Manajemen Keamanan",
            "",
            "## Cakupan",
            "Data flow · Algoritma analisis · Spesifikasi UI Tinjau Laporan Patroli",
            "",
            "## Di Luar Cakupan (v0.3)",
            "SRMD Monitoring · Profiling · Modul Input Data (detail) · Role & permission penuh",
        ])
        self._footer(slide, "Cover")

    def agenda(self) -> None:
        slide = self._blank()
        y = self._title_bar(slide, "Agenda", "Ringkasan bab presentasi")
        cols = [
            ("01  Konteks & Latar Belakang", [
                "Ringkasan eksekutif & masalah bisnis",
                "Persona pengguna & peran",
                "Arsitektur sistem",
            ]),
            ("02  Data & Integrasi", [
                "Alur data end-to-end",
                "Window analisis & fetch GPS",
                "Master data & klasifikasi tracking",
            ]),
            ("03  Business Rules & Algoritma", [
                "Exception report (berhenti, rute, kecepatan)",
                "Metrik & coverage blok rawan",
                "Penentuan shift",
            ]),
            ("04  UI / UX Requirements", [
                "User journey PIC",
                "Spesifikasi per halaman + mockup",
                "Validasi submit & workflow pasca-submit",
            ]),
            ("05  Lanjutan", [
                "Modul future scope",
                "Pertanyaan terbuka untuk IT",
            ]),
        ]
        col_w = Inches(2.35)
        x0 = MARGIN_L
        for i, (heading, items) in enumerate(cols):
            x = x0 + i * (col_w + Inches(0.15))
            self._label_chip(slide, heading.split("  ")[0], x, y, BLUE)
            self._text_block(slide, x, y + Inches(0.38), col_w, Inches(4.8), [heading.split("  ", 1)[1]] + [""] + items, size=Pt(10))
        self._footer(slide, "Agenda")

    def bullets_slide(self, title: str, subtitle: str, sections: list[tuple[str, list[str]]], section: str) -> None:
        slide = self._blank()
        y = self._title_bar(slide, title, subtitle)
        n = len(sections)
        col_w = (CONTENT_W - Inches(0.2 * (n - 1))) / n
        for i, (heading, items) in enumerate(sections):
            x = MARGIN_L + i * (col_w + Inches(0.2))
            self._label_chip(slide, heading, x, y, TEAL if i % 2 == 0 else BLUE)
            self._text_block(slide, x, y + Inches(0.38), col_w, Inches(5.2), items, size=Pt(10))
        self._footer(slide, section)

    def req_and_design(
        self,
        title: str,
        subtitle: str,
        req_heading: str,
        requirements: list[str],
        design_label: str,
        design_caption: str = "",
        section: str = "",
        design_right: bool = True,
    ) -> None:
        """Split slide: requirements left, screenshot placeholder right."""
        slide = self._blank()
        y = self._title_bar(slide, title, subtitle)
        split = Inches(5.9)
        gap = Inches(0.25)
        if design_right:
            req_x, req_w = MARGIN_L, split
            img_x, img_w = MARGIN_L + split + gap, CONTENT_W - split - gap
        else:
            img_x, img_w = MARGIN_L, CONTENT_W - split - gap
            req_x, req_w = MARGIN_L + img_w + gap, split

        self._label_chip(slide, "REQUIREMENTS", req_x, y, GREEN_LABEL)
        self._text_block(slide, req_x, y + Inches(0.35), req_w, Inches(5.3), [f"## {req_heading}", ""] + requirements, size=Pt(10))

        self._label_chip(slide, "DESIGN / MOCKUP", img_x, y, ORANGE_LABEL)
        self._screenshot_placeholder(slide, img_x, y + Inches(0.35), img_w, Inches(5.3), design_label, design_caption)
        self._footer(slide, section or title)

    def full_screenshot(self, title: str, subtitle: str, label: str, caption: str, section: str) -> None:
        slide = self._blank()
        y = self._title_bar(slide, title, subtitle)
        self._screenshot_placeholder(slide, MARGIN_L, y, CONTENT_W, Inches(5.35), label, caption)
        self._footer(slide, section)

    def closing(self) -> None:
        slide = self._blank()
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        tx = slide.shapes.add_textbox(MARGIN_L, Inches(2.8), Inches(10), Inches(1))
        p = tx.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = "Terima Kasih"
        r.font.size = Pt(40)
        r.font.bold = True
        r.font.color.rgb = WHITE
        sub = slide.shapes.add_textbox(MARGIN_L, Inches(3.9), Inches(10), Inches(0.8))
        p2 = sub.text_frame.paragraphs[0]
        r2 = p2.add_run()
        r2.text = "SmartTrack · Smart Security Dashboard  |  Dokumen lengkap: DOKUMENTASI-SMARTTRACK-v0.3.docx"
        r2.font.size = Pt(12)
        r2.font.color.rgb = RGBColor(0xBB, 0xCC, 0xEE)
        self._footer(slide, "Closing")

    def save(self) -> None:
        self.prs.save(str(OUT))
        print(f"Saved: {OUT} ({len(self.prs.slides)} slides)")


def build() -> None:
    d = Deck()

    # ── Cover & Agenda ────────────────────────────────────────────────
    d.cover()
    d.agenda()

    # ── SECTION 01: Konteks ───────────────────────────────────────────
    d._section_divider("01", "Konteks & Latar Belakang", "Masalah bisnis, solusi, dan pengguna sistem")

    d.bullets_slide(
        "Ringkasan Eksekutif",
        "SmartTrack — modul patroli dalam Smart Security Dashboard",
        [
            ("Masalah", [
                "Pelaporan patroli minim visibilitas: sulit verifikasi rute, durasi berhenti, cakupan blok rawan.",
                "Data GPS belum terintegrasi ke proses review harian PIC & manajemen.",
                "Tidak ada deteksi otomatis anomaly (berhenti lama, keluar rute, speeding).",
            ]),
            ("Solusi", [
                "Platform terpusat: inventaris → jadwal rute → fetch GPS → analisis → review PIC.",
                "Laporan terstruktur dengan exception report & metrik kinerja per shift.",
                "Workflow audit trail: PIC submit → Manajer review → RC (fase berikutnya).",
            ]),
            ("Nilai Bisnis", [
                "Akuntabilitas patroli di area berisiko tinggi.",
                "Data-driven decision untuk alokasi tim & perbaikan rute.",
                "Standardisasi laporan across estate.",
            ]),
        ],
        "01 · Konteks",
    )

    d.bullets_slide(
        "Persona & Peran Pengguna",
        "Siapa yang berinteraksi dengan sistem dan apa tanggung jawabnya",
        [
            ("Admin / Input Data", [
                "Mengelola inventaris kendaraan & IMEI GPS.",
                "Mengatur rute bank, blok rawan, jadwal shift.",
                "Plot rute ke kendaraan per tanggal (3–9 rute/hari).",
            ]),
            ("PIC", [
                "Meninjau laporan harian di Tinjau Laporan Patroli.",
                "Memberi komentar wajib pada setiap exception.",
                "Submit laporan + ringkasan & rencana tindakan.",
            ]),
            ("Manajer & RC", [
                "Review laporan di Hasil Laporan Patroli.",
                "Memberi komentar & arahan tindak lanjut.",
                "Menerima notifikasi khusus untuk laporan Urgent.",
            ]),
        ],
        "01 · Konteks",
    )

    d.req_and_design(
        "Arsitektur Sistem",
        "Empat lapisan yang saling berinteraksi",
        "Komponen",
        [
            "Lapisan 1 — Master Data: inventaris, rute, blok rawan, jadwal.",
            "Lapisan 2 — GPS: perangkat IMEI → server GPS eksternal (real-time).",
            "Lapisan 3 — Engine Analisis: klasifikasi, deteksi event, statistik.",
            "Lapisan 4 — Workflow: antrian PIC → submit → review manajemen.",
            "",
            "Integrasi kunci: fetch harian dari server GPS eksternal ke Smart Security Dashboard.",
        ],
        "Diagram Arsitektur 4 Lapis",
        "Sisipkan diagram: Master Data → GPS → Engine → Workflow",
        "01 · Arsitektur",
    )

    # ── SECTION 02: Data ──────────────────────────────────────────────
    d._section_divider("02", "Data & Integrasi", "Alur data, window waktu, master data, klasifikasi GPS")

    d.req_and_design(
        "Alur Data End-to-End",
        "Dari input master data hingga laporan disubmit",
        "Tahapan Proses",
        [
            "A. Admin input inventaris kendaraan + IMEI GPS.",
            "B. Plot rute ke kendaraan per tanggal & shift (1–3 rute/shift).",
            "C. Patroli berjalan → GPS kirim data ke server eksternal.",
            "D. Fetch jam 08:00 — data window 07:00 (D) s.d. 07:00 (D+1).",
            "E. Klasifikasi tracking: Ada (≥1 km) / Tidak Ada (<1 km).",
            "F. Engine analisis → event & statistik.",
            "G. Laporan muncul di antrian Tinjau Laporan Patroli.",
            "H. PIC review → submit → Hasil Laporan Patroli → Manajemen.",
        ],
        "Diagram Alur Data A → H",
        "Flowchart end-to-end — dapat digambar di Visio",
        "02 · Data Flow",
    )

    d.bullets_slide(
        "Window Analisis vs Waktu Fetch",
        "Dua konsep waktu yang sengaja dipisahkan",
        [
            ("Window Analisis", [
                "Satu hari patroli = 07:00 (D) → 07:00 (D+1).",
                "Bukan kalender 00:00–23:59.",
                "Contoh: laporan 14 Feb = data 14 Feb 07:00 – 15 Feb 07:00.",
                "Tanggal laporan = hari D (awal window).",
            ]),
            ("Waktu Fetch", [
                "Setiap pagi jam 08:00 dari server GPS.",
                "Buffer 1 jam setelah window ditutup.",
                "Alasan: hindari data pending dari Shift 3 (berakhir 07:00).",
            ]),
            ("Shift Operasional", [
                "Shift 1: 07:00 – 15:00",
                "Shift 2: 15:00 – 23:00",
                "Shift 3: 23:00 – 07:00 (lintas hari)",
                "Pengecualian: 06:42–06:59 → berhenti tidak dianalisis.",
            ]),
        ],
        "02 · Waktu",
    )

    d.bullets_slide(
        "Master Data — Input Requirements",
        "Data yang harus tersedia sebelum patroli dianalisis",
        [
            ("Inventaris", [
                "Estate, tipe kendaraan (Mobil/Motor), plat nomor.",
                "IMEI GPS terpasang, status kendaraan & perangkat.",
                "Relasi: 1 kendaraan ↔ 1 IMEI aktif per periode.",
            ]),
            ("Rute Bank", [
                "ID, nama rute, estate, geometri GeoJSON.",
                "Panjang rute dihitung otomatis dari polyline.",
            ]),
            ("Penugasan Harian", [
                "Tanggal, kendaraan, shift, 1–3 rute per shift.",
                "Total: 3–9 rute baseline per kendaraan per hari.",
            ]),
            ("Blok Rawan", [
                "Kode blok, estate, polygon GeoJSON.",
                "Flag rawan / non-rawan.",
            ]),
        ],
        "02 · Master Data",
    )

    d.req_and_design(
        "Klasifikasi Ketersediaan Tracking",
        "Keputusan pertama sebelum analisis dimulai",
        "Business Rules",
        [
            "Hitung total jarak GPS dalam window 07:00–07:00.",
            "",
            "Total < 1 km → Tidak Tersedia → Trouble Report.",
            "Total ≥ 1 km → Tersedia → Laporan Normal.",
            "",
            "Tidak ada data sama sekali → Tidak Tersedia.",
            "",
            "Catatan: meski Tersedia, PIC bisa tekan INTERVENSI",
            "(contoh: motor rusak, GPS masih rekam ke bengkel).",
        ],
        "Diagram Klasifikasi Laporan",
        "Flowchart: GPS data → jarak? → Normal / Trouble / Intervensi",
        "02 · Klasifikasi",
    )

    # ── SECTION 03: Algoritma ─────────────────────────────────────────
    d._section_divider("03", "Business Rules & Algoritma", "Logika deteksi event dan perhitungan metrik")

    d.bullets_slide(
        "Parameter Sistem (Konfigurasi)",
        "Nilai default — dapat disesuaikan per estate",
        [
            ("Waktu & Tracking", [
                "Window harian: 07:00 – 07:00 (D+1)",
                "Fetch: 08:00 harian",
                "Threshold tidak ada tracking: 1 km/hari",
                "Exclude berhenti: 06:42 – 06:59",
            ]),
            ("Deteksi Event", [
                "Berhenti: radius 50 m, durasi 30 min",
                "Kecepatan: > 50 km/jam",
                "Keluar rute: > 1 km dari baseline",
                "On-route buffer: 25 m (20 m jalan + 3–5 m GPS)",
            ]),
            ("Blok Rawan", [
                "Min. jarak lewat blok: 600 m",
                "Gap antar kunjungan: 23 menit",
                "Min. kata ringkasan: 10",
            ]),
        ],
        "03 · Parameter",
    )

    d.req_and_design(
        "Exception: Berhenti > 30 Menit",
        "REQ-EXC-001",
        "Detection Logic",
        [
            "Trigger: GPS tidak bergerak signifikan > 30 menit, ATAU",
            "tetap dalam radius ≤ 50 m dari satu titik > 30 menit.",
            "",
            "Output: waktu mulai/selesai, koordinat, blok terdekat, shift.",
            "Shift ditentukan dari waktu MULAI event.",
            "",
            "Exclude: rentang 06:42–06:59 tidak dihitung.",
            "",
            "PIC wajib komentar: Standby / Istirahat / Operasional / Lainnya.",
            "Sub-kategori wajib untuk Standby (di Pos, Kantor, Workshop).",
            "",
            "Enhancement: geofencing auto-reason (pondok, pos).",
        ],
        "Exception Panel — Berhenti",
        "Screenshot: accordion Berhenti expanded + dropdown komentar",
        "03 · Berhenti",
    )

    d.req_and_design(
        "Exception: Keluar Rute per Shift",
        "REQ-EXC-002",
        "Detection Logic",
        [
            "On-route: jarak ke polyline baseline ≤ 25 m.",
            "Off-route: jarak ke titik terdekat rute > 1 km (configurable).",
            "",
            "Hanya hitung saat kendaraan bergerak (speed > 0).",
            "Minimal 5 titik GPS berurutan untuk 1 event.",
            "",
            "Output: waktu mulai/kembali, track deviasi, total jarak, shift.",
            "",
            "PIC komentar: Mendadak / Jalur / Operasional / Lainnya.",
            "Sub wajib untuk Operasional (BBM, koordinasi, akses blok).",
        ],
        "Exception Panel — Keluar Rute",
        "Screenshot: accordion Keluar Rute + segment di peta",
        "03 · Keluar Rute",
    )

    d.req_and_design(
        "Exception: Kecepatan > 50 km/jam",
        "REQ-EXC-003",
        "Detection Logic",
        [
            "Hitung kecepatan antar 2 titik GPS berurutan.",
            "Trigger: kecepatan > 50 km/jam.",
            "Akhiri grup setelah 2 titik berturut-turut ≤ 50 km/jam (anti-noise).",
            "",
            "Output: waktu, kecepatan maks, total jarak, track, shift.",
            "",
            "PIC komentar: Darurat / Operasional / Lainnya.",
            "Sub wajib untuk Operasional (mengejar target, respon insiden).",
        ],
        "Exception Panel — Kecepatan",
        "Screenshot: accordion Kecepatan + track pink di peta",
        "03 · Kecepatan",
    )

    d.bullets_slide(
        "Metrik Analisis & Coverage",
        "Perhitungan yang mengisi panel Analisis Patroli & Statistik",
        [
            ("Donut Charts", [
                "Waktu bergerak vs berhenti (% durasi).",
                "Coverage blok rawan vs non-rawan (% jarak).",
                "Blok rawan dicek: belum / 1× / 2× / >2×.",
            ]),
            ("Panel Statistik", [
                "Jarak rute dijadwalkan vs jarak aktual.",
                "Keluar rute (km) · Sesuai rute · Rute tercover · % cover.",
            ]),
            ("Blok Rawan — Aturan", [
                "Kunjungan valid: jarak dalam blok ≥ 600 m.",
                "Gap antar kunjungan sama: ≥ 23 menit.",
                "Tabel per shift: waktu bergerak, kecepatan rata-rata, % cover.",
            ]),
            ("Kecepatan", [
                "Rata-rata, maksimal, minimal per shift & per hari.",
            ]),
        ],
        "03 · Metrik",
    )

    d.req_and_design(
        "Penentuan Shift pada Exception",
        "REQ-SHIFT-001",
        "Business Rule",
        [
            "Shift ditentukan dari WAKTU MULAI event — bukan selesai.",
            "",
            "Shift 1: start 07:00 – 14:59",
            "Shift 2: start 15:00 – 22:59",
            "Shift 3: start 23:00 – 06:59 (D+1)",
            "",
            "Contoh:",
            "• Berhenti 14:55–15:27 → Shift 1",
            "• Speeding 06:42–06:43 → Shift 3",
            "• Keluar rute 00:07–00:11 → Shift 3",
            "",
            "Visualisasi peta:",
            "Track S1 merah · S2 biru · S3 cyan",
            "Baseline: garis transparan per shift",
        ],
        "Peta — Warna per Shift",
        "Screenshot: peta dengan legend shift + baseline",
        "03 · Shift",
    )

    # ── SECTION 04: UI/UX ─────────────────────────────────────────────
    d._section_divider("04", "UI / UX Requirements", "Spesifikasi tampilan dan alur pengguna PIC")

    d.req_and_design(
        "User Journey — PIC",
        "Alur utama dari antrian hingga submit",
        "Flow",
        [
            "1. Buka Tinjau Laporan Patroli (antrian).",
            "2. Lihat stat cards: belum ditinjau, pelacakan, tipe kendaraan.",
            "3. Pilih baris laporan → cek ketersediaan tracking.",
            "",
            "Jika Tidak Tersedia:",
            "→ Form kendala (GPS rusak / kendaraan rusak) → Kirim.",
            "",
            "Jika Tersedia:",
            "→ Review peta + exception → komentar semua event",
            "→ Isi ringkasan & rencana tindakan (min. 10 kata)",
            "→ Set prioritas Normal/Urgent → KIRIM.",
            "",
            "Opsi INTERVENSI kapan saja dari laporan normal.",
        ],
        "User Journey Diagram",
        "Flowchart: Antrian → Normal/Trouble → Submit → Hasil Laporan",
        "04 · Journey",
    )

    d.full_screenshot(
        "UI-001: Halaman Indeks — Tinjau Laporan Patroli",
        "Antrian laporan belum disubmit oleh PIC",
        "Halaman Indeks /review-patrol",
        "Stat cards (belum ditinjau · pelacakan · kendaraan) + tabel 7 kolom",
        "04 · UI Indeks",
    )

    d.req_and_design(
        "UI-002: Form Laporan Kendala (Trouble Report)",
        "Ditampilkan saat tracking Tidak Tersedia atau PIC menekan INTERVENSI",
        "Requirements",
        [
            "Metadata: Estate · Tipe kendaraan · Plat · Tanggal.",
            "",
            "Form:",
            "• Jenis kerusakan (radio): Motor/Unit Rusak ATAU GPS/WiFi Rusak.",
            "• Keterangan (textarea, wajib).",
            "",
            "Aksi: Batal | Kirim Laporan.",
            "",
            "Tidak perlu: exception comment, analisis, ringkasan 10 kata.",
            "",
            "Setelah kirim: laporan keluar dari antrian review.",
        ],
        "Form Laporan Kendala",
        "Screenshot: /review-patrol/[id]/kendala",
        "04 · UI Kendala",
    )

    d.full_screenshot(
        "UI-003: Laporan Hasil Patroli — Overview",
        "Halaman review utama saat tracking Tersedia",
        "Laporan Hasil Patroli /review-patrol/[id]/hasil",
        "Header (metadata + INTERVENSI + KIRIM + prioritas) + peta + exception panel",
        "04 · UI Hasil",
    )

    d.req_and_design(
        "UI-004: Peta Interaktif & Playback",
        "REQ-UI-MAP-001",
        "Map Requirements",
        [
            "Default load: polygon blok estate + highlight blok rawan (merah).",
            "Layer baseline rute per shift (transparan, tebal).",
            "Layer track aktual GPS per shift (solid, warna shift).",
            "",
            "Marker event: berhenti (oranye), keluar rute, speeding (pink).",
            "Event dikomentari → marker hijau.",
            "",
            "Playback: Play · Pause · Resume · progress bar 0–100%.",
            "Animasi pergerakan kendaraan sepanjang track.",
            "",
            "Klik event di panel → fokus ke lokasi di peta.",
        ],
        "Peta + Playback Controls",
        "Screenshot: peta dengan legend + playback bar",
        "04 · UI Peta",
    )

    d.full_screenshot(
        "UI-005: Analisis Patroli & Statistik",
        "Panel bawah halaman review — donut charts + metrik jarak",
        "Analisis Patroli + Statistik + Tabel per Shift",
        "3 donut charts · 6 baris statistik · tabel 3 shift dengan % cover berwarna",
        "04 · UI Analisis",
    )

    d.req_and_design(
        "UI-006: Penugasan Tim & Ringkasan",
        "REQ-UI-FORM-001",
        "Form Requirements",
        [
            "Penugasan tim (grid 3×3 dropdown):",
            "Manajer (pre-filled, locked) · Kanitpam · PIC",
            "Danru Shift 1–3 · Anggota Shift 1–3",
            "",
            "Ringkasan Patroli (textarea, wajib, min. 10 kata).",
            "Rencana Tindakan (textarea, wajib, min. 10 kata).",
            "",
            "View-only setelah submit (kecuali komentar Manajer).",
        ],
        "Form Penugasan + Ringkasan",
        "Screenshot: bagian bawah halaman review",
        "04 · UI Form",
    )

    d.bullets_slide(
        "Validasi Submit & Pasca-Submit",
        "Kondisi tombol KIRIM aktif dan alur setelah submit",
        [
            ("Syarat Submit", [
                "✓ Semua exception dikomentari (berhenti, rute, kecepatan).",
                "✓ Ringkasan patroli ≥ 10 kata.",
                "✓ Rencana tindakan ≥ 10 kata.",
                "✓ Penugasan tim terisi (PIC minimal — TBD).",
                "✓ Prioritas dipilih: Normal / Urgent.",
            ]),
            ("Setelah Submit", [
                "Laporan pindah ke Hasil Laporan Patroli.",
                "Hilang dari antrian Tinjau Laporan Patroli.",
                "Normal → notifikasi Manajer.",
                "Urgent → notifikasi Manajer + Manajemen.",
            ]),
            ("Review Lanjutan", [
                "Manajer: komentar di laporan disubmit.",
                "RC: kolom komentar (fase berikutnya).",
            ]),
        ],
        "04 · Submit",
    )

    # ── SECTION 05: Lanjutan ────────────────────────────────────────────
    d._section_divider("05", "Future Scope & Open Items", "Modul lanjutan dan pertanyaan untuk tim IT")

    d.bullets_slide(
        "Modul di Luar Cakupan v0.3",
        "Akan didokumentasikan setelah mockup tersedia",
        [
            ("Dashboard", [
                "SRMD Monitoring — mockup static SPA tersedia.",
                "Profiling — overview man power & infrastruktur.",
            ]),
            ("Input Data (6 sub-modul)", [
                "Tag Rute · GPS & Kendaraan · Anggota Tim",
                "Blok Rawan · Rute Bank · Jadwal",
            ]),
            ("Laporan", [
                "Hasil Laporan Patroli — indeks + detail view-only.",
                "Komentar Manajer & RC.",
            ]),
        ],
        "05 · Future",
    )

    d.bullets_slide(
        "Pertanyaan Terbuka untuk Tim IT",
        "Keputusan yang memerlukan konfirmasi sebelum development",
        [
            ("Integrasi", [
                "Platform & protokol server GPS eksternal?",
                "PostGIS vs GeoJSON untuk penyimpanan geometri?",
                "Retensi data GPS mentah?",
            ]),
            ("Workflow", [
                "Field penugasan tim mana yang wajib sebelum submit?",
                "Mekanisme notifikasi Urgent (push / email / WA)?",
                "Apakah trouble report perlu approval manajer?",
            ]),
            ("Dokumen", [
                "Referensi lengkap: DOKUMENTASI-SMARTTRACK-v0.3.docx",
                "Mockup UI: environment dev /review-patrol",
            ]),
        ],
        "05 · Open Items",
    )

    d.closing()
    d.save()


if __name__ == "__main__":
    build()
